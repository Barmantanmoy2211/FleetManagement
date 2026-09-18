from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook

from app.models import FuelType, LeaseOwnershipType, VehicleStatus, VehicleType

VEHICLE_IMPORT_HEADERS: list[tuple[str, str]] = [
    ("vehicleName", "vehicleName"),
    ("make", "make"),
    ("yearOfManufacture", "yearOfManufacture"),
    ("model", "model"),
    ("vin", "vin"),
    ("color", "color"),
    ("registrationNumber", "registrationNumber"),
    ("dotNumber", "dotNumber"),
    ("leaseOwnershipType", "leaseOwnershipType"),
    ("vehicleStatus", "vehicleStatus"),
    ("policyNumber", "policyNumber"),
    ("coveredUnderPolicy", "coveredUnderPolicy"),
    ("cargoType", "cargoType"),
    ("weightLbs", "weightLbs"),
    ("vehicleType", "vehicleType"),
    ("fuelType", "fuelType"),
    ("vehicleSubtype", "vehicleSubtype"),
]

HEADER_LABELS = [label for label, _ in VEHICLE_IMPORT_HEADERS]

SAMPLE_ROW = [
    "Heavy Truck- 1XKAD89X97J109196",
    "Kenworth",
    "2007",
    "T6 series",
    "",
    "",
    "MH12AB1234",
    "",
    "",
    "Active",
    "CS SPS 5483052-01",
    "yes",
    "",
    "",
    "Heavy Truck",
    "",
    "",
]


def build_vehicle_import_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Vehicles"
    ws.append(HEADER_LABELS)
    ws.append(SAMPLE_ROW)
    instructions = wb.create_sheet("Instructions")
    instructions.append(["Field", "Notes"])
    instructions.append(["vehicleName", "Required"])
    instructions.append(["make / model / yearOfManufacture", "Required"])
    instructions.append(["vin / dotNumber", "Optional for now"])
    instructions.append(["vehicleStatus", "Active maps to AVAILABLE, or use AVAILABLE"])
    instructions.append(["vehicleType", "TRUCK, VAN, CAR, BUS, BIKE, OTHER, or Heavy Truck"])
    instructions.append(["fuelType", "DIESEL, PETROL, CNG, ELECTRIC, HYBRID, OTHER"])
    instructions.append(["coveredUnderPolicy", "yes or no"])
    instructions.append(["tenant", "Set by upload scope — not a column"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _normalize_cell(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _parse_bool(value: str | None) -> bool:
    if not value:
        return False
    return value.strip().lower() in ("1", "true", "yes", "y")


def _parse_year(value: str | None) -> int | None:
    if not value:
        return None
    try:
        year = int(float(value))
    except ValueError as exc:
        raise ValueError(f"Invalid year '{value}'") from exc
    if year < 1980 or year > 2100:
        raise ValueError(f"Year out of range: {year}")
    return year


def _parse_vehicle_type(value: str | None) -> VehicleType:
    if not value:
        return VehicleType.TRUCK
    stripped = value.strip()
    lower = stripped.lower()
    friendly = {
        "heavy truck": VehicleType.TRUCK,
        "truck": VehicleType.TRUCK,
        "tractor": VehicleType.TRUCK,
        "van": VehicleType.VAN,
        "car": VehicleType.CAR,
        "bus": VehicleType.BUS,
        "bike": VehicleType.BIKE,
        "motorcycle": VehicleType.BIKE,
        "other": VehicleType.OTHER,
    }
    if lower in friendly:
        return friendly[lower]
    normalized = stripped.upper().replace(" ", "_")
    aliases = {
        "HEAVY_TRUCK": VehicleType.TRUCK,
        "HEAVYTRUCK": VehicleType.TRUCK,
        "TRACTOR": VehicleType.TRUCK,
    }
    if normalized in aliases:
        return aliases[normalized]
    for vt in VehicleType:
        if vt.value == normalized or vt.name == normalized:
            return vt
    return VehicleType.OTHER


def _parse_fuel_type(value: str | None) -> FuelType:
    if not value:
        return FuelType.DIESEL
    stripped = value.strip()
    lower = stripped.lower()
    friendly = {
        "diesel": FuelType.DIESEL,
        "petrol": FuelType.PETROL,
        "gasoline": FuelType.PETROL,
        "gas": FuelType.PETROL,
        "cng": FuelType.CNG,
        "electric": FuelType.ELECTRIC,
        "ev": FuelType.ELECTRIC,
        "hybrid": FuelType.HYBRID,
        "other": FuelType.OTHER,
    }
    if lower in friendly:
        return friendly[lower]
    normalized = stripped.upper()
    for ft in FuelType:
        if ft.value == normalized:
            return ft
    return FuelType.OTHER


def _parse_vehicle_status(value: str | None) -> VehicleStatus:
    if not value:
        return VehicleStatus.AVAILABLE
    normalized = value.strip().upper()
    if normalized in ("ACTIVE", "AVAILABLE"):
        return VehicleStatus.AVAILABLE
    try:
        return VehicleStatus(normalized)
    except ValueError as exc:
        raise ValueError(f"Invalid vehicleStatus '{value}'") from exc


def _parse_lease_type(value: str | None) -> LeaseOwnershipType | None:
    if not value:
        return None
    normalized = value.strip().upper()
    for lt in LeaseOwnershipType:
        if lt.value == normalized:
            return lt
    # Demo / legacy sheets often put vehicle category here — ignore unknown values.
    return None


def _parse_weight(value: str | None) -> float | None:
    if not value:
        return None
    try:
        w = float(value)
    except ValueError as exc:
        raise ValueError(f"Invalid weightLbs '{value}'") from exc
    if w < 0:
        raise ValueError("weightLbs must be >= 0")
    return w


def parse_vehicle_import_rows(
    file_bytes: bytes,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return [], [{"row": 1, "message": "Empty workbook"}]

    header_index: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        label = _normalize_cell(cell)
        if label and label in HEADER_LABELS:
            header_index[label] = idx

    if "vehicleName" not in header_index:
        return [], [{"row": 1, "message": "Missing required column: vehicleName"}]

    valid: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []

    for row_num, row in enumerate(rows_iter, start=2):
        if not row or all(_normalize_cell(c) is None for c in row):
            continue

        def get_col(label: str) -> str | None:
            idx = header_index.get(label)
            if idx is None or idx >= len(row):
                return None
            return _normalize_cell(row[idx])

        vehicle_name = get_col("vehicleName")
        if not vehicle_name:
            continue

        make = get_col("make")
        model = get_col("model")
        if not make or not model:
            errors.append({"row": row_num, "message": "make and model are required"})
            continue

        try:
            year = _parse_year(get_col("yearOfManufacture"))
            if year is None:
                errors.append({"row": row_num, "message": "yearOfManufacture is required"})
                continue

            valid.append(
                {
                    "vehicleName": vehicle_name,
                    "make": make,
                    "model": model,
                    "year": year,
                    "vin": get_col("vin") or None,
                    "color": get_col("color"),
                    "registrationNumber": get_col("registrationNumber"),
                    "dotNumber": get_col("dotNumber") or None,
                    "leaseOwnershipType": _parse_lease_type(get_col("leaseOwnershipType")),
                    "status": _parse_vehicle_status(get_col("vehicleStatus")),
                    "policyNumber": get_col("policyNumber"),
                    "coveredUnderPolicy": _parse_bool(get_col("coveredUnderPolicy")),
                    "cargoType": get_col("cargoType"),
                    "weightLbs": _parse_weight(get_col("weightLbs")),
                    "vehicleType": _parse_vehicle_type(get_col("vehicleType")),
                    "fuelType": _parse_fuel_type(get_col("fuelType")),
                    "vehicleSubtype": get_col("vehicleSubtype"),
                }
            )
        except ValueError as exc:
            errors.append({"row": row_num, "message": str(exc)})

    return valid, errors
