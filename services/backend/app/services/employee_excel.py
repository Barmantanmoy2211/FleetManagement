from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook

from app.models import EmployeePersona, EmployeeStatus, Gender

# Header row must match template download and import parser.
EMPLOYEE_IMPORT_HEADERS: list[tuple[str, str]] = [
    ("name", "name"),
    ("employeeCode", "employeeCode"),
    ("persona", "persona"),
    ("dateOfBirth", "dateOfBirth"),
    ("gender", "gender"),
    ("status", "status"),
    ("isDriver", "isDriver"),
    ("hireDate", "hireDate"),
    ("homeAddress", "homeAddress"),
    ("email", "email"),
    ("phone", "phone"),
    ("primaryContact", "primaryContact"),
    ("city", "city"),
    ("state", "state"),
    ("zipCode", "zipCode"),
    ("country", "country"),
    ("emergencyContactName", "emergencyContactName"),
    ("emergencyContactAddress", "emergencyContactAddress"),
    ("employmentType", "employmentType"),
    ("employmentStatus", "employmentStatus"),
    ("experience", "experience"),
    ("dailyHoursWorked", "dailyHoursWorked"),
    ("companyDriverId", "companyDriverId"),
    ("department", "department"),
    ("jobRole", "jobRole"),
    ("driverManagerEmail", "driverManagerEmail"),
]

HEADER_LABELS = [label for label, _ in EMPLOYEE_IMPORT_HEADERS]

SAMPLE_ROW = [
    "Jane Doe",
    "1001",
    EmployeePersona.DRIVER.value,
    "1990-05-15",
    Gender.FEMALE.value,
    EmployeeStatus.ACTIVE.value,
    "yes",
    "2024-01-10",
    "123 Main St",
    "jane.doe@example.com",
    "+1-555-0100",
    "no",
    "Austin",
    "TX",
    "78701",
    "USA",
    "John Doe",
    "456 Oak Ave",
    "Full-time",
    "ACTIVE",
    "5 years",
    "8",
    "DRV-1001",
    "Operations",
    "Line haul",
    "",
]


def build_employee_import_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(HEADER_LABELS)
    ws.append(SAMPLE_ROW)
    instructions = wb.create_sheet("Instructions")
    instructions.append(["Field", "Notes"])
    instructions.append(["name", "Required"])
    instructions.append(["persona", "Fleet Admin | Fleet Manager | Driver"])
    instructions.append(["driverManagerEmail", "Fleet Manager user email (Driver persona only)"])
    instructions.append(["dateOfBirth / hireDate", "YYYY-MM-DD"])
    instructions.append(["isDriver / primaryContact", "yes or no"])
    instructions.append(["status", "ACTIVE | INACTIVE | ON_LEAVE | TERMINATED"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _normalize_cell(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _parse_bool(value: str | None) -> bool:
    if not value:
        return False
    return value.strip().lower() in ("1", "true", "yes", "y")


def _parse_persona(value: str | None) -> EmployeePersona | None:
    if not value:
        return None
    normalized = value.strip().lower().replace(" ", "")
    mapping = {
        "fleetadmin": EmployeePersona.FLEET_ADMIN,
        "locationhead": EmployeePersona.LOCATION_HEAD,
        "fleetmanager": EmployeePersona.FLEET_MANAGER,
        "driver": EmployeePersona.DRIVER,
    }
    if normalized in mapping:
        return mapping[normalized]
    for persona in EmployeePersona:
        if persona.value.lower() == value.strip().lower():
            return persona
    raise ValueError(
        f"Invalid persona '{value}' — use Fleet Admin, Location Head, Fleet Manager, or Driver"
    )


def _parse_gender(value: str | None) -> Gender | None:
    if not value:
        return None
    for g in Gender:
        if g.value.lower() == value.strip().lower():
            return g
    raise ValueError(f"Invalid gender '{value}'")


def _parse_status(value: str | None) -> EmployeeStatus:
    if not value:
        return EmployeeStatus.ACTIVE
    try:
        return EmployeeStatus(value.strip().upper())
    except ValueError as exc:
        raise ValueError(f"Invalid status '{value}'") from exc


def parse_employee_import_rows(file_bytes: bytes) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Returns (valid_rows, errors) where errors are {row, message}."""
    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return [], [{"row": 1, "message": "Empty workbook"}]

    header_index: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        label = _normalize_cell(cell)
        if label and label in HEADER_LABELS:
            header_index[label] = idx

    if "name" not in header_index:
        return [], [{"row": 1, "message": "Missing required column: name"}]

    valid: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []

    for row_num, row in enumerate(rows_iter, start=2):
        if not row or all(_normalize_cell(c) is None for c in row):
            continue

        def get_col(label: str) -> str | None:
            idx = header_index.get(label)
            if idx is None or idx >= len(row):
                return None
            return _normalize_cell(row[idx])

        name = get_col("name")
        if not name:
            continue

        try:
            persona_raw = get_col("persona")
            persona = _parse_persona(persona_raw)
            is_driver = _parse_bool(get_col("isDriver"))
            if persona == EmployeePersona.DRIVER:
                is_driver = True

            daily_raw = get_col("dailyHoursWorked")
            daily_hours = float(daily_raw) if daily_raw else None

            gender_str = None
            gender_raw = get_col("gender")
            if gender_raw:
                gender_str = _parse_gender(gender_raw).value

            valid.append(
                {
                    "name": name,
                    "employeeCode": get_col("employeeCode"),
                    "persona": persona.value if persona else None,
                    "dateOfBirth": get_col("dateOfBirth"),
                    "gender": gender_str,
                    "status": _parse_status(get_col("status")).value,
                    "isDriver": is_driver,
                    "hireDate": get_col("hireDate"),
                    "homeAddress": get_col("homeAddress"),
                    "email": get_col("email"),
                    "phone": get_col("phone"),
                    "primaryContact": _parse_bool(get_col("primaryContact")),
                    "city": get_col("city"),
                    "state": get_col("state"),
                    "zipCode": get_col("zipCode"),
                    "country": get_col("country"),
                    "emergencyContactName": get_col("emergencyContactName"),
                    "emergencyContactAddress": get_col("emergencyContactAddress"),
                    "employmentType": get_col("employmentType"),
                    "employmentStatus": get_col("employmentStatus"),
                    "experience": get_col("experience"),
                    "dailyHoursWorked": daily_hours,
                    "companyDriverId": get_col("companyDriverId"),
                    "department": get_col("department"),
                    "jobRole": get_col("jobRole"),
                    "driverManagerEmail": get_col("driverManagerEmail"),
                }
            )
        except ValueError as exc:
            errors.append({"row": row_num, "message": str(exc)})

    return valid, errors
